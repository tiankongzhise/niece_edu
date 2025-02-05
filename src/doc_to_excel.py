
import os
import docx
from openpyxl import Workbook

def extract_tables_from_word_to_excel(word_file_path, excel_file_path):
    # 打开Word文档
    doc = docx.Document(word_file_path)

    # 创建一个新的Excel工作簿
    wb = Workbook()

    # 遍历Word文档中的每个表格
    for table_index, table in enumerate(doc.tables, start=1):
        # 如果是第一个表格，使用默认的工作表；否则，创建一个新的工作表
        if table_index == 1:
            ws = wb.active
        else:
            ws = wb.create_sheet(f'Table {table_index}')

        # 遍历表格的每一行
        for row in table.rows:
            # 存储当前行的数据
            row_data = []
            # 遍历行中的每个单元格
            for cell in row.cells:
                # 将单元格的文本添加到行数据列表中
                row_data.append(cell.text)
            # 将行数据添加到Excel工作表中
            ws.append(row_data)

    # 保存Excel文件
    wb.save(excel_file_path)
    print(f"表格已成功从 {word_file_path} 提取并保存到 {excel_file_path}")

if __name__ == '__main__':
    word_file_path = r'./data/doc/'
    excel_file_path = r'./data/2025/'
    os.makedirs(os.path.dirname(excel_file_path), exist_ok=True)
    all_files = os.listdir(word_file_path)
    doc_files = [file for file in all_files if file.endswith(('.docx'))]
    for doc in doc_files:
        extract_tables_from_word_to_excel(word_file_path+doc, f'{excel_file_path}{os.path.splitext(doc)[0]}.xlsx')

